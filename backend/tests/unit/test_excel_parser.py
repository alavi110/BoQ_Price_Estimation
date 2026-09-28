"""
Unit tests for Excel parser
"""
import pytest
from src.services.boq_parser import BoQParser, parse_boq_excel, ParsedBoQItem

def test_parse_valid_excel(tmp_path):
    """Test parsing a valid BoQ Excel file"""
    # Create a test Excel file
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Header row
    ws.append(["کد", "شرح", "واحد", "مبلغ", "تعداد"])
    
    # Data rows
    ws.append(["101001", "بتنی ساده", "متر مکعب", 5000000, 10])
    ws.append(["101002", "بتنی آرمه", "متر مکعب", 7000000, 5])
    ws.append(["201001", "کابل مسی 3x150", "متر", 2500000, 100])
    
    file_path = tmp_path / "test_boq.xlsx"
    wb.save(file_path)
    
    # Parse
    parser = BoQParser()
    result = parser.parse(str(file_path))
    
    assert len(result.items) == 3
    assert result.items[0].code == "101001"
    assert result.items[0].description_fa == "بتنی ساده"
    assert result.items[0].chapter_code == "101"
    assert result.items[2].chapter_code == "201"
    assert result.warnings == []


def test_parse_missing_header(tmp_path):
    """Test parsing Excel without proper header"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # No header row - just data
    ws.append(["101001", "بتنی ساده", "متر مکعب", 5000000, 10])
    
    file_path = tmp_path / "test_no_header.xlsx"
    wb.save(file_path)
    
    parser = BoQParser()
    with pytest.raises(ValueError, match="Could not find header row"):
        parser.parse(str(file_path))


def test_chapter_extraction():
    """Test chapter code extraction from item codes"""
    parser = BoQParser()
    
    assert parser._extract_chapter_code("101001") == "101"
    assert parser._extract_chapter_code("201005") == "201"
    assert parser._extract_chapter_code("301") == "301"
    assert parser._extract_chapter_code("ABC123") == "000"
    assert parser._extract_chapter_code("") == "000"


def test_persian_text_handling(tmp_path):
    """Test Persian RTL text handling"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    
    ws.append(["کد", "شرح", "واحد", "مبلغ", "تعداد"])
    ws.append(["101001", "لوله پالایشگاه جنوب پارس", "متر", 15000000, 50])
    
    file_path = tmp_path / "test_persian.xlsx"
    wb.save(file_path)
    
    parser = BoQParser()
    result = parser.parse(str(file_path))
    
    assert len(result.items) == 1
    assert "پالایشگاه" in result.items[0].description_fa
    assert "جنوب" in result.items[0].description_fa
    assert "پارس" in result.items[0].description_fa